import io
import re
import unicodedata
from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


st.set_page_config(
    page_title="Wordwall Results Manager",
    page_icon="📚",
    layout="wide"
)

PATRON = re.compile(r"^\s*(\d+[A-Za-z]?)\s+(.+?)\s*$")

CORRECTAS = {
    "✔", "✓", "true", "correct", "correcto",
    "1", "yes", "si", "sí"
}

INCORRECTAS = {
    "✖", "✘", "✗", "✕", "❌", "false",
    "incorrect", "incorrecto", "0", "no"
}


def normalizar(texto):
    texto = unicodedata.normalize("NFKD", str(texto))
    caracteres = "".join(
        c for c in texto
        if not unicodedata.combining(c)
    )
    return " ".join(caracteres.casefold().split())


def extraer_intentos(tabla, archivo):
    filas = tabla.fillna("").astype(str).values.tolist()

    actividad = next(
        (fila[0].strip() for fila in filas if fila and fila[0].strip()),
        Path(archivo).stem
    )

    intentos = []
    alumno_actual = None
    aciertos = 0
    preguntas = 0

    def guardar():
        if alumno_actual is not None and preguntas > 0:
            grupo, alumno = alumno_actual
            intentos.append({
                "Grupo": grupo,
                "Alumno": alumno,
                "Actividad": actividad,
                "Calificación": 100 * aciertos / preguntas
            })

    for fila in filas[1:]:
        a = fila[0].strip() if len(fila) > 0 else ""
        b = fila[1].strip() if len(fila) > 1 else ""
        c = fila[2].strip() if len(fila) > 2 else ""

        coincidencia = PATRON.fullmatch(a) if not b and not c else None

        if coincidencia:
            guardar()
            alumno_actual = (
                coincidencia.group(1),
                coincidencia.group(2).strip()
            )
            aciertos = 0
            preguntas = 0
            continue

        if alumno_actual is None:
            continue

        marca = normalizar(c)

        if marca in CORRECTAS or marca in INCORRECTAS:
            preguntas += 1
            if marca in CORRECTAS:
                aciertos += 1

    guardar()
    return actividad, intentos


def leer_archivo(archivo):
    datos = archivo.getvalue()

    if archivo.name.lower().endswith(".csv"):
        for codificacion in ("utf-8-sig", "cp1252", "latin-1"):
            try:
                tabla = pd.read_csv(
                    io.BytesIO(datos),
                    header=None,
                    dtype=str,
                    sep=None,
                    engine="python",
                    encoding=codificacion
                )
                return extraer_intentos(tabla, archivo.name)
            except UnicodeError:
                continue

        raise ValueError("No se pudo leer el archivo CSV")

    tabla = pd.read_excel(
        io.BytesIO(datos),
        sheet_name="Results by student",
        header=None,
        dtype=str,
        engine="openpyxl"
    )

    return extraer_intentos(tabla, archivo.name)


def crear_tablas(registros, actividades):
    datos = pd.DataFrame(registros)

    datos["Clave"] = datos["Alumno"].map(normalizar)

    datos["Intento"] = datos.groupby(
        ["Grupo", "Clave", "Actividad"],
        sort=False
    ).cumcount() + 1

    primeros = datos[datos["Intento"] <= 3]

    notas = primeros.groupby(
        ["Grupo", "Clave", "Actividad"],
        as_index=False
    ).agg(
        Calificación=("Calificación", "mean"),
        Alumno=("Alumno", "first")
    )

    resultado = {}

    grupos = sorted(
        datos["Grupo"].unique(),
        key=lambda x: (
            not str(x).isdigit(),
            str(x).zfill(10)
        )
    )

    for grupo in grupos:
        parte = notas[notas["Grupo"] == grupo]

        tabla = parte.pivot_table(
            index="Clave",
            columns="Actividad",
            values="Calificación",
            aggfunc="mean"
        )

        tabla = tabla.reindex(columns=actividades)

        nombres = parte.groupby("Clave")["Alumno"].first()

        tabla.insert(
            0,
            "Alumno",
            nombres.reindex(tabla.index)
        )

        tabla[actividades] = tabla[actividades].fillna(0)

        tabla["Promedio mensual"] = (
            tabla[actividades].mean(axis=1).round(2)
        )

        tabla[actividades] = tabla[actividades].round(2)

        resultado[grupo] = (
            tabla.reset_index(drop=True)
            .sort_values(
                "Alumno",
                key=lambda columna: columna.map(normalizar)
            )
        )

    return resultado


def generar_excel(tablas):
    salida = io.BytesIO()

    with pd.ExcelWriter(salida, engine="openpyxl") as escritor:
        for grupo, tabla in tablas.items():
            nombre_hoja = ("Grupo " + str(grupo))[:31]

            tabla.to_excel(
                escritor,
                sheet_name=nombre_hoja,
                index=False
            )

            hoja = escritor.sheets[nombre_hoja]
            hoja.freeze_panes = "B2"
            hoja.auto_filter.ref = hoja.dimensions

            for celda in hoja[1]:
                celda.font = Font(
                    bold=True,
                    color="FFFFFF"
                )
                celda.fill = PatternFill(
                    "solid",
                    fgColor="174A69"
                )
                celda.alignment = Alignment(
                    wrap_text=True
                )

            for columna in hoja.columns:
                ancho = min(
                    55,
                    max(
                        15,
                        max(
                            len(str(c.value or ""))
                            for c in columna
                        ) + 2
                    )
                )

                letra = get_column_letter(
                    columna[0].column
                )

                hoja.column_dimensions[letra].width = ancho

    return salida.getvalue()


st.title("📚 Wordwall Results Manager")

st.write(
    "Unifica las calificaciones de Wordwall "
    "por grupo y actividad."
)

st.info(
    "Cada actividad es una columna. "
    "Se promedian hasta los tres primeros intentos "
    "y las actividades pendientes valen cero."
)

mes = st.text_input(
    "Título del reporte",
    value="Reporte de Wordwall",
    placeholder="Escribe el título de tu reporte"
)

archivos = st.file_uploader(
    "Sube los archivos Excel o CSV de Wordwall",
    type=["xlsx", "csv"],
    accept_multiple_files=True
)

st.caption(
    "El orden de los intentos será el orden de aparición "
    "en el archivo. No subas el mismo archivo dos veces."
)

if archivos:
    registros = []
    actividades = []

    for archivo in archivos:
        try:
            actividad, intentos = leer_archivo(archivo)

            if not intentos:
                st.warning(
                    f"No se encontraron resultados en {archivo.name}"
                )
                continue

            if actividad in actividades:
                st.warning(
                    f"Actividad repetida: {actividad}. "
                    "Revisa que no hayas subido dos veces "
                    "la misma actividad."
                )
                continue

            actividades.append(actividad)
            registros.extend(intentos)

        except Exception as error:
            st.error(
                f"Error al leer {archivo.name}: {error}"
            )

    if registros:
        tablas = crear_tablas(
            registros,
            actividades
        )

        st.success(
            f"Se encontraron {len(tablas)} grupos "
            f"y {len(actividades)} actividades."
        )

        grupo = st.selectbox(
            "Selecciona un grupo para ver sus calificaciones",
            list(tablas)
        )

        st.dataframe(
            tablas[grupo],
            hide_index=True,
            use_container_width=True
        )

        nombre_mes = re.sub(
            r"[^A-Za-z0-9_-]+",
            "_",
            mes
        ).strip("_") or "Mensual"

        st.download_button(
            "⬇️ Descargar Excel por grupos",
            data=generar_excel(tablas),
            file_name=f"Wordwall_{nombre_mes}.xlsx",
            mime=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    else:
        st.error(
            "No se encontraron resultados válidos."
        )
